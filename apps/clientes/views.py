import io

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from .importer import importar_clientes, importar_contratos_locacao, importar_contratos_venda
from .models import Cliente, ContratoLocacao, ContratoVenda, ImportLog

_IMPORTERS = {
    'clientes': importar_clientes,
    'vendas':   importar_contratos_venda,
    'locacoes': importar_contratos_locacao,
}
_LABELS = {
    'clientes': 'Clientes',
    'vendas':   'Contratos de Venda',
    'locacoes': 'Contratos de Locação',
}


def _fmt_brl(v):
    v = v or 0
    formatted = f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return f'R$ {formatted}'


def _filtrar_clientes(request):
    qs = Cliente.objects.annotate(
        n_vendas=Count('vendas', distinct=True),
        n_locacoes=Count('locacoes', distinct=True),
    )

    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(nome__icontains=q) | Q(codigo__icontains=q) | Q(cpf_cnpj__icontains=q))

    tipo_cliente = request.GET.get('tipo_cliente', '').strip()
    if tipo_cliente:
        qs = qs.filter(tipo_cliente=tipo_cliente)

    relacao = request.GET.get('relacao', '').strip()
    if relacao == 'com_venda':
        qs = qs.filter(n_vendas__gt=0)
    elif relacao == 'com_locacao':
        qs = qs.filter(n_locacoes__gt=0)
    elif relacao == 'ambos':
        qs = qs.filter(n_vendas__gt=0, n_locacoes__gt=0)
    elif relacao == 'nenhuma':
        qs = qs.filter(n_vendas=0, n_locacoes=0)

    return qs.order_by('nome')


@login_required
def lista(request):
    clientes = _filtrar_clientes(request)
    tipos_cliente = list(
        Cliente.objects.exclude(tipo_cliente='').values_list('tipo_cliente', flat=True).distinct().order_by('tipo_cliente')
    )
    context = {
        'clientes':      clientes,
        'n_clientes':    clientes.count(),
        'tipos_cliente': tipos_cliente,
        'filtros': {
            'q':            request.GET.get('q', ''),
            'tipo_cliente': request.GET.get('tipo_cliente', ''),
            'relacao':      request.GET.get('relacao', ''),
        },
    }
    return render(request, 'clientes/lista.html', context)


@login_required
def detalhe(request, codigo):
    cliente = get_object_or_404(Cliente, codigo=codigo)
    context = {
        'cliente':  cliente,
        'vendas':   cliente.vendas.all(),
        'locacoes': cliente.locacoes.all(),
    }
    return render(request, 'clientes/detalhe.html', context)


@login_required
def importar(request):
    if request.method == 'POST':
        tipo = request.POST.get('tipo', '')
        arquivo = request.FILES.get('arquivo')
        if not tipo or tipo not in _IMPORTERS:
            messages.error(request, 'Tipo de arquivo inválido.')
            return redirect('clientes:importar')
        if not arquivo:
            messages.error(request, 'Nenhum arquivo selecionado.')
            return redirect('clientes:importar')
        try:
            n = _IMPORTERS[tipo](arquivo, arquivo.name)
            messages.success(request, f'{_LABELS[tipo]} importado — {n} registros.')
        except Exception as e:
            messages.error(request, f'Erro ao importar {_LABELS[tipo]}: {e}')
        return redirect('clientes:importar')

    ultimos = {}
    for log in ImportLog.objects.all():
        if log.tipo not in ultimos:
            ultimos[log.tipo] = log

    arquivos = []
    for tipo, label in _LABELS.items():
        log = ultimos.get(tipo)
        arquivos.append({
            'tipo':      tipo,
            'label':     label,
            'ultimo':    log.importado_em.strftime('%d/%m/%Y %H:%M') if log else None,
            'registros': log.total_registros if log else 0,
            'arquivo':   log.nome_arquivo if log else '—',
        })

    context = {
        'arquivos':     arquivos,
        'n_clientes':   Cliente.objects.count(),
        'n_vendas':     ContratoVenda.objects.count(),
        'n_locacoes':   ContratoLocacao.objects.count(),
    }
    return render(request, 'clientes/importar.html', context)


# ── Exportação ───────────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill('solid', fgColor='1a1a2e')
_HEADER_FONT = Font(color='FFFFFF', bold=True)
_ALT_FILL    = PatternFill('solid', fgColor='f8f9fa')


@login_required
def exportar(request):
    fmt = request.GET.get('format', 'xlsx')
    clientes = _filtrar_clientes(request)
    if fmt == 'pdf':
        return _exportar_pdf(clientes)
    return _exportar_xlsx(clientes)


def _exportar_xlsx(clientes):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Clientes'

    ws.merge_cells('A1:G1')
    ws['A1'] = 'Grupo Cota — Lista de Clientes'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')

    headers = ['CÓDIGO', 'NOME', 'TIPO DE CLIENTE', 'TELEFONE', 'EMAIL', 'VENDAS', 'LOCAÇÕES']
    widths  = [10, 38, 20, 18, 32, 10, 10]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = w

    for i, c in enumerate(clientes):
        row = 4 + i
        fill = _ALT_FILL if i % 2 == 0 else PatternFill()
        data = [c.codigo, c.nome, c.tipo_cliente, c.telefone, c.email, c.n_vendas, c.n_locacoes]
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = fill

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="clientes.xlsx"'
    return resp


def _exportar_pdf(clientes):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    story = [
        Paragraph('Grupo Cota — Lista de Clientes',
                  ParagraphStyle('title', parent=styles['Heading1'], fontSize=14, spaceAfter=4)),
        Paragraph(f'Total de Clientes: {clientes.count()}',
                  ParagraphStyle('sub', parent=styles['Normal'], fontSize=10, spaceAfter=12)),
    ]

    header = ['CÓDIGO', 'NOME', 'TIPO DE CLIENTE', 'TELEFONE', 'EMAIL', 'VENDAS', 'LOCAÇÕES']
    data = [header]
    for c in clientes:
        data.append([c.codigo, c.nome, c.tipo_cliente, c.telefone, c.email, str(c.n_vendas), str(c.n_locacoes)])

    t = Table(data, colWidths=[1.5*cm, 6*cm, 3*cm, 3*cm, 5.5*cm, 2*cm, 2*cm], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND',    (0, 0), (-1, 0), colors.HexColor('#1a1a2e')),
        ('TEXTCOLOR',     (0, 0), (-1, 0), colors.white),
        ('FONTNAME',      (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE',      (0, 0), (-1, 0), 8),
        ('FONTSIZE',      (0, 1), (-1, -1), 7),
        ('ROWBACKGROUNDS',(0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
        ('GRID',          (0, 0), (-1, -1), 0.5, colors.HexColor('#dee2e6')),
        ('TOPPADDING',    (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    resp = HttpResponse(buf, content_type='application/pdf')
    resp['Content-Disposition'] = 'attachment; filename="clientes.pdf"'
    return resp
